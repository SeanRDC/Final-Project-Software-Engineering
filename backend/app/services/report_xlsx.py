# Builds the Excel copy of a clinic report: the university letterhead, then each table under a red header.

from collections.abc import Callable, Sequence
from datetime import date, datetime
from io import BytesIO
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.clock import to_clinic_time
from app.schemas.support import ReportSummary

LOGO = Path(__file__).resolve().parent.parent / "assets" / "hau-logo.png"

# The app's colours (frontend/src/index.css): the clinic red and quiet greys.
RED = "7A1C2A"
WHITE = "FFFFFF"
BAND = "F7F2F3"
LINE = "D9D4D5"
MUTED = "6B6466"

COLUMNS = 5
WIDTHS = [44, 16, 16, 16, 16]

HEADER_FILL = PatternFill("solid", fgColor=RED)
BAND_FILL = PatternFill("solid", fgColor=BAND)
EDGE = Side(style="thin", color=LINE)
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)


def _readable(label: str) -> str:
    """"medicine_request" -> "Medicine request", as the screens show it."""
    if label == "other":
        return "Others"
    words = label.replace("_", " ").strip()
    return words[:1].upper() + words[1:]


def _month(label: str) -> str:
    """"2026-10" -> "October 2026"."""
    try:
        return datetime.strptime(label, "%Y-%m").strftime("%B %Y")
    except ValueError:
        return label


def _long_date(day: date) -> str:
    return f"{day:%B} {day.day}, {day.year}"


class _Sheet:
    """Writes the report top to bottom, one table after another."""

    def __init__(self, sheet: Worksheet) -> None:
        self.sheet = sheet
        self.row = 1

    def letterhead(self, title: str, period: str, generated: str) -> None:
        lines = [
            ("HOLY ANGEL UNIVERSITY", Font(name="Calibri", size=16, bold=True, color=RED)),
            ("University Clinic", Font(name="Calibri", size=12, bold=True)),
            (title, Font(name="Calibri", size=13, bold=True)),
            (period, Font(name="Calibri", size=11)),
            (generated, Font(name="Calibri", size=9, italic=True, color=MUTED)),
        ]
        for text, font in lines:
            self.sheet.merge_cells(
                start_row=self.row, start_column=1, end_row=self.row, end_column=COLUMNS
            )
            cell = self.sheet.cell(self.row, 1, text)
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            self.sheet.row_dimensions[self.row].height = 20
            self.row += 1
        if LOGO.exists():
            logo = Image(str(LOGO))
            logo.width = logo.height = 84
            self.sheet.add_image(logo, "A1")
        # A red rule under the letterhead.
        for column in range(1, COLUMNS + 1):
            self.sheet.cell(self.row, column).border = Border(bottom=Side(style="medium", color=RED))
        self.sheet.row_dimensions[self.row].height = 6
        self.row += 2

    def table(self, headings: Sequence[str], rows: Sequence[Sequence[object]]) -> None:
        # Number columns are right-aligned, headings included, so the figures line up.
        numeric = {
            column for column in range(1, len(headings) + 1)
            if rows and isinstance(rows[0][column - 1], int | float)
        }
        for column, heading in enumerate(headings, start=1):
            cell = self.sheet.cell(self.row, column, heading)
            cell.fill = HEADER_FILL
            cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="right" if column in numeric else "left", vertical="center"
            )
        self.sheet.row_dimensions[self.row].height = 20
        self.row += 1
        if not rows:
            cell = self.sheet.cell(self.row, 1, "None in this period")
            cell.font = Font(name="Calibri", size=10, italic=True, color=MUTED)
            cell.border = BOX
            self.row += 2
            return
        for index, values in enumerate(rows):
            lines = 1
            for column, value in enumerate(values, start=1):
                cell = self.sheet.cell(self.row, column, value)
                cell.font = Font(name="Calibri", size=11)
                cell.border = BOX
                if index % 2 == 1:
                    cell.fill = BAND_FILL
                if isinstance(value, int | float):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                else:
                    # Long names wrap inside their column instead of being cut off.
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    lines = max(lines, ceil(len(str(value)) / (WIDTHS[column - 1] * 1.05)))
            if lines > 1:
                self.sheet.row_dimensions[self.row].height = 15 * lines
            self.row += 1
        self.row += 1


def build(
    summary: ReportSummary,
    title: str,
    generated_by: str | None,
    generated_at: datetime,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.sheet_view.showGridLines = False
    for column, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    local = to_clinic_time(generated_at)
    stamp = f"{_long_date(local.date())}, {local:%I:%M %p}".replace(", 0", ", ")
    period = f"{_long_date(summary.period_start)} to {_long_date(summary.period_end)}"
    writer = _Sheet(sheet)
    writer.letterhead(
        title,
        period,
        f"Generated {stamp}" + (f" by {generated_by}" if generated_by else ""),
    )

    writer.table(["Overview", "Total"], [
        ["Total visits", summary.total_visits],
        ["Unique patients", summary.unique_patients],
        ["Referrals", summary.referrals],
        ["Guardian notifications", summary.guardian_notifications],
    ])

    def counts(heading: str, items, label: Callable[[str], str] = _readable) -> None:
        writer.table([heading, "Count"], [[label(item.label), item.count] for item in items])

    counts("Visits by patient type", summary.visits_by_patient_type)
    counts("Visits by type of request", summary.visits_by_type)
    if summary.visits_by_type_and_patient_type:
        writer.table(
            ["Requests by type of patient", "Students", "Employees", "Others", "Total"],
            [
                [_readable(item.label), item.students, item.employees, item.others,
                 item.students + item.employees + item.others]
                for item in summary.visits_by_type_and_patient_type
            ],
        )
    counts("Visits by department", summary.visits_by_department, label=str)
    counts("Visits by month", summary.visits_by_month, label=_month)
    counts("Visits by outcome", summary.visits_by_disposition)
    counts("After admission", summary.admission_releases)
    counts("Top complaints", summary.top_complaints)
    counts("Appointments by status", summary.appointments_by_status)
    writer.table(
        ["Medicines released", "Quantity", "Unit"],
        [[m.medicine_name, m.quantity_released, m.unit] for m in summary.medicines_released],
    )
    writer.table(
        ["Frequent visitors", "ID number", "Department", "Visits"],
        [[v.full_name, v.id_number, v.department or "", v.visit_count]
         for v in summary.frequent_visitors],
    )

    # Ready to print: one page wide, as many pages tall as it takes, centred.
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.5
    sheet.oddFooter.center.text = "Page &P of &N"
    sheet.oddFooter.left.text = title

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
